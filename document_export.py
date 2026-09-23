"""DOCX export shared by Streamlit and local batch verification."""
from __future__ import annotations
import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pipeline import questions, stamp


def make_docx(lines: list[dict], result: dict, actions: list[dict],
              title: str = "Протокол совещания",
              diarized_lines: list[dict] | None = None) -> bytes:
    if not str(result.get("summary") or "").strip():
        raise ValueError("Нельзя экспортировать протокол без краткого содержания.")
    doc = Document()
    # The bundled default template can carry a decorative title border.
    for root in (doc.styles.element, doc.element):
        for border in list(root.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.65)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    doc.add_paragraph(title, "Title")
    doc.add_paragraph("Черновик по записи. Проверьте факты, ответственных, сроки и отмеченные фрагменты по источнику.")
    doc.add_heading("Краткое содержание", 1)
    doc.add_paragraph(result.get("summary", ""))
    if result.get("review_notes"):
        doc.add_heading("Факты для проверки", 1)
        for note in result["review_notes"]:
            doc.add_paragraph(str(note), "List Bullet")
    if result.get("decisions"):
        doc.add_heading("Решения", 1)
        for decision in result["decisions"]:
            doc.add_paragraph(str(decision), "List Bullet")
    doc.add_heading("Поручения", 1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    widths = [3.15, 1.45, 1.30, 1.30]
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for cell, title, width in zip(table.rows[0].cells, ["Задача по ASR", "Ответственный", "Срок", "Источник"], widths):
        cell.width = Inches(width)
        cell.text = title
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i, action in enumerate(actions, 1):
        cells = table.add_row().cells
        values = [f"{i}. {action['task']}", action.get("owner") or "Не указан",
                  action.get("deadline") or "Не указан",
                  ", ".join(str(n) for n in action.get("evidence_ids", []))]
        for cell, value, width in zip(cells, values, widths):
            cell.text = str(value)
            cell.width = Inches(width)
        no_split = OxmlElement("w:cantSplit")
        table.rows[-1]._tr.get_or_add_trPr().append(no_split)
    doc.add_heading("Основания поручений", 1)
    for i, action in enumerate(actions, 1):
        p = doc.add_paragraph()
        p.add_run(f"Поручение {i}. ").bold = True
        p.add_run(action.get("evidence", ""))
        for note in action.get("review_notes", []):
            doc.add_paragraph("Проверить: " + str(note))
    issues = questions(actions)
    if issues:
        doc.add_heading("Требует уточнения", 1)
        for issue in issues:
            doc.add_paragraph(issue, "List Bullet")
    doc.add_page_break()
    doc.add_heading("ASR транскрипт", 1)
    doc.add_paragraph("Номера строк соответствуют источникам в таблице. Метка «проверить говорящего» означает неопределённую атрибуцию.")
    for line in lines:
        p = doc.add_paragraph()
        end = stamp(line.get("end", line.get("start", 0)))
        label = f"[{line['id']}] {line['time']}-{end} {line['speaker']}"
        if line.get("needs_review"):
            label += " — проверить говорящего"
        p.add_run(label + ": ").bold = True
        p.add_run(line["text"])
    if diarized_lines is not None:
        doc.add_page_break()
        doc.add_heading("Диаризация по словам", 1)
        doc.add_paragraph("Сегменты с неопределённым говорящим явно помечены; поручения ссылаются на номера в ASR транскрипте выше.")
        for line in diarized_lines:
            p = doc.add_paragraph()
            end = stamp(line.get("end", line.get("start", 0)))
            label = f'[{line["id"]}] {line["time"]}-{end} {line["speaker"]}'
            if line.get("needs_review"):
                label += " — проверить говорящего"
            p.add_run(label + ": ").bold = True
            p.add_run(line["text"])
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
