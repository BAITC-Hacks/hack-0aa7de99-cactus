"""Сборка протокола в DOCX по извлечённым данным."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json, datetime

def build(data: dict, out: str = "Протокол_авто.docx",
          org: str = 'АО «Самрук-Қазына Ондеу»', date: str = "2026-09-22",
          transcript: str | None = None):
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)

    h = d.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run("Протокол совещания"); r.bold = True; r.font.size = Pt(16)
    sub = d.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"{org}\n{date}").italic = True

    if data.get("topics"):
        p = d.add_paragraph(); p.add_run("Повестка: ").bold = True
        p.add_run("; ".join(data["topics"]))

    if data.get("participants"):
        d.add_heading("Участники", level=2)
        for x in data["participants"]:
            role = f" — {x['role']}" if x.get("role") else ""
            d.add_paragraph(f"{x['name']}{role}", style="List Bullet")

    d.add_heading("Поручения", level=2)
    t = d.add_table(rows=1, cols=5); t.style = "Table Grid"
    for i, c in enumerate(["№", "Поручение", "Ответственный", "Срок", "Основание"]):
        cell = t.rows[0].cells[i]; cell.text = ""
        cell.paragraphs[0].add_run(c).bold = True
    for i, task in enumerate(data.get("tasks", []), 1):
        cells = t.add_row().cells
        due = task.get("due_date") or task.get("due_raw") or "—"
        cells[0].text = str(i)
        cells[1].text = task.get("task", "")
        cells[2].text = task.get("assignee", "—")
        cells[3].text = str(due)
        q = (task.get("quote") or "")[:160]
        cells[4].paragraphs[0].add_run(f"«{q}»").italic = True

    for w, col in zip([Cm(1), Cm(6.5), Cm(3.2), Cm(2.3), Cm(5)], t.columns):
        for c in col.cells: c.width = w

    if data.get("summary"):
        d.add_heading("Саммари", level=2)
        d.add_paragraph(data["summary"])

    if transcript:
        d.add_page_break(); d.add_heading("Приложение: стенограмма", level=2)
        for line in transcript.split("\n"):
            if line.strip(): d.add_paragraph(line.strip())

    f = d.add_paragraph()
    f.add_run(f"\nСформировано автоматически {datetime.date.today()}. "
              "Обработка выполнена локально, без передачи данных во внешние сервисы."
              ).italic = True
    d.save(out)
    return out

if __name__ == "__main__":
    data = json.load(open("tasks_out.json"))
    tr = json.load(open("stt_out.json"))["text"]
    print("сохранено:", build(data, transcript=tr))
