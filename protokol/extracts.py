"""Персональные выписки из протокола — по одной на ответственного.

Опциональный пункт кейса: «автоматическая рассылка выдержки из протокола
ответственным лицам». Саму отправку намеренно не делаем: письмо из
закрытого контура наружу — это ровно то, что кейс запрещает. Готовим
файлы и текст письма, отправку выполняет почтовая система заказчика.
"""
import os, re, datetime
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

URGENCY_ORDER = {"высокая": 0, "средняя": 1, "низкая": 2}


def _safe(name: str) -> str:
    return re.sub(r"[^\w\s-]", "", name).strip().replace(" ", "_") or "без_имени"


def by_assignee(tasks: list[dict]) -> dict[str, list[dict]]:
    out: dict[str, list[dict]] = {}
    for t in tasks:
        who = (t.get("assignee") or "не назначен").strip()
        out.setdefault(who, []).append(t)
    for who in out:
        out[who].sort(key=lambda t: (URGENCY_ORDER.get(t.get("urgency"), 1),
                                     t.get("due_date") or "9999"))
    return out


def letter_text(who: str, tasks: list[dict], date: str, org: str) -> str:
    """текст письма — его остаётся только вставить в почтовый клиент"""
    lines = [f"Тема: Поручения по итогам совещания {date}", "",
             f"{who}, добрый день.", "",
             f"По итогам совещания {org} от {date} на вас оформлено "
             f"поручений: {len(tasks)}.", ""]
    for i, t in enumerate(tasks, 1):
        due = t.get("due_date") or t.get("due_raw") or "срок не указан"
        urg = t.get("urgency")
        mark = " (срочно)" if urg == "высокая" else ""
        lines.append(f"{i}. {t.get('task')}{mark}")
        lines.append(f"   Срок: {due}")
        if t.get("quote"):
            lines.append(f"   Основание: «{t['quote'][:150]}»")
        lines.append("")
    lines += ["Протокол во вложении.", "",
              "Письмо сформировано автоматически системой протоколирования."]
    return "\n".join(lines)


def build(data: dict, out_dir: str = "выписки",
          date: str = "", org: str = 'АО «Самрук-Қазына Ондеу»') -> list[str]:
    """Кладёт в папку DOCX-выписку и текст письма на каждого ответственного."""
    os.makedirs(out_dir, exist_ok=True)
    made = []
    for who, tasks in by_assignee(data.get("tasks", [])).items():
        d = Document()
        d.styles["Normal"].font.name = "Times New Roman"
        d.styles["Normal"].font.size = Pt(11)

        h = d.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = h.add_run("Выписка из протокола совещания"); r.bold = True
        r.font.size = Pt(14)
        sub = d.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.add_run(f"{org} · {date}").italic = True

        p = d.add_paragraph(); p.add_run("Ответственный: ").bold = True
        p.add_run(who)

        t = d.add_table(rows=1, cols=5); t.style = "Table Grid"
        for i, c in enumerate(["№", "Поручение", "Срок", "Срочность", "Основание"]):
            t.rows[0].cells[i].paragraphs[0].add_run(c).bold = True
        for i, task in enumerate(tasks, 1):
            c = t.add_row().cells
            c[0].text = str(i)
            c[1].text = task.get("task", "")
            c[2].text = str(task.get("due_date") or task.get("due_raw") or "—")
            c[3].text = task.get("urgency") or "—"
            c[4].paragraphs[0].add_run(f"«{(task.get('quote') or '')[:140]}»").italic = True

        f = d.add_paragraph()
        f.add_run(f"\nСформировано автоматически {datetime.date.today()}. "
                  "Обработка выполнена локально.").italic = True

        base = os.path.join(out_dir, _safe(who))
        d.save(base + ".docx")
        open(base + "_письмо.txt", "w").write(letter_text(who, tasks, date, org))
        made.append(base + ".docx")
    return made
