from __future__ import annotations

import io
import tempfile
from pathlib import Path

import streamlit as st
from docx import Document

from pipeline import diarize, extract, lines_from_segments, parse_manual, questions, transcribe


def make_docx(lines: list[dict], result: dict, actions: list[dict]) -> bytes:
    doc = Document()
    doc.add_heading("Протокол совещания", 0)
    doc.add_paragraph("Черновик: проверьте ответственных, сроки и исходные реплики перед утверждением.")
    doc.add_heading("Краткое содержание", 1)
    doc.add_paragraph(result["summary"])
    doc.add_heading("Решения", 1)
    for decision in result["decisions"]:
        doc.add_paragraph(decision, style="List Bullet")
    doc.add_heading("Поручения", 1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, title in zip(table.rows[0].cells, ["Задача", "Ответственный", "Срок", "Основание"]):
        cell.text = title
    for action in actions:
        cells = table.add_row().cells
        for cell, key in zip(cells, ["task", "owner", "deadline", "evidence"]):
            cell.text = str(action.get(key) or "Не указан")
    doc.add_heading("Требует уточнения", 1)
    for issue in questions(actions):
        doc.add_paragraph(issue, style="List Bullet")
    doc.add_heading("Транскрипт", 1)
    for line in lines:
        doc.add_paragraph(f'{line["time"]} {line["speaker"]}: {line["text"]}')
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


st.set_page_config(page_title="Meeting Execution Intelligence", layout="wide")
st.title("Протокол совещания")
st.caption("Локальная обработка аудио · говорящие · поручения · вопросы на уточнение")

audio = st.file_uploader("Загрузите запись MP3 или WAV", type=["mp3", "wav", "m4a"])
manual = st.text_area("Или вставьте размеченный транскрипт для проверки извлечения",
                      placeholder="Асхат: Нурлан, подготовьте отчёт до пятницы.\nНурлан: Принято.",
                      height=100)

if st.button("1. Распознать запись / принять транскрипт", type="primary"):
    try:
        if manual.strip():
            st.session_state.lines = parse_manual(manual)
            st.session_state.diarization_status = "Ручной транскрипт: ASR и диаризация не проверялись."
        elif audio:
            with tempfile.TemporaryDirectory() as folder:
                path = Path(folder) / f"meeting{Path(audio.name).suffix.lower()}"
                path.write_bytes(audio.getvalue())
                with st.spinner("Локальное распознавание…"):
                    segments = transcribe(path)
                try:
                    with st.spinner("Локальная диаризация…"):
                        segments = diarize(path, segments)
                    st.session_state.diarization_status = "Диаризация выполнена."
                except Exception as error:
                    st.session_state.diarization_status = f"Диаризация не выполнена: {error}"
                st.session_state.lines = lines_from_segments(segments)
        else:
            st.warning("Добавьте запись или транскрипт.")
        st.session_state.pop("result", None)
    except Exception as error:
        st.error(str(error))

if "lines" in st.session_state:
    lines = st.session_state.lines
    st.info(st.session_state.diarization_status)
    st.subheader("Транскрипт и говорящие")
    speakers = sorted({line["speaker"] for line in lines if line["speaker"] != "Не определён"})
    if speakers:
        renames = {speaker: st.text_input(f"Имя для {speaker}", value=speaker,
                                          key=f"speaker_{speaker}") for speaker in speakers}
        for line in lines:
            line["speaker"] = renames.get(line["speaker"], line["speaker"])
    for line in lines:
        st.write(f'**{line["time"]} · {line["speaker"]}** — {line["text"]}')
    if st.button("2. Извлечь поручения локально"):
        try:
            with st.spinner("Анализ текста локальной моделью…"):
                st.session_state.result = extract(lines)
                st.session_state.edited_actions = st.session_state.result["actions"]
        except Exception as error:
            st.error(str(error))

if "result" in st.session_state:
    result = st.session_state.result
    st.subheader("Краткое содержание")
    st.write(result["summary"])
    st.subheader("Решения")
    for decision in result["decisions"]:
        st.write("•", decision)
    st.subheader("Поручения — проверьте и исправьте")
    actions = st.data_editor(st.session_state.edited_actions, num_rows="dynamic",
                             column_config={"evidence_ids": None}, key="actions_editor",
                             use_container_width=True)
    st.subheader("Что уточнить до завершения совещания?")
    issues = questions(actions)
    if issues:
        for issue in issues:
            st.warning(issue)
    else:
        st.success("У всех обнаруженных поручений указан исполнитель и срок.")
    st.download_button("Скачать протокол DOCX", make_docx(st.session_state.lines, result, actions),
                       file_name="meeting_protocol.docx",
                       mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
